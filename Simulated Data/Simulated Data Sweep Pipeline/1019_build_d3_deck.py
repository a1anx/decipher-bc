"""Build the d3 sigma-sweep deck: v-space grids per sigma, then the compare_d3 metric plots.

Usage (repo root, .venv):
    .venv/bin/python "Simulated Data/Simulated Data Sweep Pipeline/1019_build_d3_deck.py"
    .venv/bin/python "Simulated Data/Simulated Data Sweep Pipeline/1019_build_d3_deck.py" --check

Each cell is the v-space PNG of one run (aspect kept, fitted to the cell) with a caption
`rho <.3f> · r² <.3f>` from the sweep ledger. A missing PNG or ledger row raises; no blank cells.
After the sigma slides, one slide per COMPARE_GROUPS entry lays out those compare_d3 `*_vs_sigma`
PNGs in a row. Those plots pool every sim seed and decipherseed, so their titles name the pooled
seeds (read from the ledgers for the deck's arms and sigmas) rather than the v-space sim seed. Every PNG in the compare dir must belong to a group, so a new plot can't be dropped.
`--check` reopens the deck and verifies slide count, titles, picture count, picture bytes
(sha256 against the source PNG) and that every picture lies inside the slide.

Subsets (added 2026-10-19 for the `sweep1019_d3ext` extension): `--ledger` may be repeated (no
run may be in two ledgers), and `--models`, `--sigmas`, `--decipher-seeds` set the columns, the
slides and the rows. E.g. the 4-arm sigma 7.5-11 deck, 6 decipherseed rows:
    ... 1019_build_d3_deck.py --ledger <1019>/sweep1019_d3_sweep_log.csv \
        --ledger <1019>/sweep1019_d3ext_sweep_log.csv --compare-dir <1019>/compare_d3_ext \
        --out <1019>/decks/1019_vspace_d3ext_4arms_sigma7p5-11_seed3_bif.pptx \
        --models model5 model4r model5r model7 --sigmas 7.5 8 9 10 11 --decipher-seeds 1 2 3 4 5 6
Decks go in the sweep's `<MMDD>/decks/` folder, never in `Claude Files/writeups/`.
"""

import argparse
import csv
import hashlib
import os
from typing import Dict, List, Optional, Sequence, Tuple, Union

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

_HERE = os.path.dirname(os.path.abspath(__file__))
_REPO = os.path.abspath(os.path.join(_HERE, "..", ".."))
_OUT_1019 = os.path.join(
    _REPO, "Simulated Data", "Sweeps and Results", "shift_sigma_sweep_bifurcation", "1019"
)
DEFAULT_LEDGER = os.path.join(_OUT_1019, "sweep1019_d3_sweep_log.csv")
DEFAULT_FIGS_DIR = os.path.join(_OUT_1019, "figs_decipher_models3")
DEFAULT_COMPARE_DIR = os.path.join(_OUT_1019, "compare_d3")
DEFAULT_OUT = os.path.join(_OUT_1019, "decks", "1019_vspace_d3_5arms_sigma_sweep_seed3_bif.pptx")

SIGMAS = [0.1, 0.5, 1.0, 2.0, 5.0, 7.5, 10.0]
MODELS = ["native", "model5", "model4r", "model5r", "model7"]
DECIPHER_SEEDS = [1, 2, 3]
# (slide title, compare_d3 PNGs left to right)
COMPARE_GROUPS = [
    (
        "Trajectory accuracy: rho",
        ["rho_vs_sigma.png", "rho_plus_vs_sigma.png", "rho_minus_vs_sigma.png"],
    ),
    ("Reconstruction: r²", ["r2_overall_vs_sigma.png", "r2_per_gene_median_vs_sigma.png"]),
    (
        "Latent structure: branch vs batch ASW",
        ["branch_asw_vs_sigma.png", "batch_asw_v_vs_sigma.png"],
    ),
    ("Training length: epochs", ["best_epoch_vs_sigma.png", "n_epochs_vs_sigma.png"]),
]
SLIDE_W_IN, SLIDE_H_IN = 13.333, 7.5

# Layout in inches
TITLE_TOP, TITLE_H = 0.15, 0.55
HEADER_TOP, HEADER_H = 0.75, 0.35
ROW_LABEL_W = 1.25
MARGIN = 0.25
CAPTION_H = 0.28
GAP = 0.08

Box = Tuple[float, float, float, float]  # left, top, width, height (inches)
Ledger = Dict[Tuple[str, float, int, int], Tuple[float, float]]
Groups = Sequence[Tuple[str, Sequence[str]]]


# ---- pure parts -------------------------------------------------------------------------


def slide_title(sigma: float, seed: int) -> str:
    return f"σ = {sigma:g} | sim seed {seed} | bif, gene-mode"


def png_path(figs_dir: str, model: str, sigma: float, seed: int, decipher_seed: int) -> str:
    """Read-only twin of 1008_simulated_data_pipeline_splitfix._fig_path for kind="vspace",
    bif, gene-mode (the original also makedirs, which a reader must not do)."""
    arm = "native_genes" if model == "native" else model
    name = f"vspace_sigma{sigma:g}_seed{seed}_{model}_bif_genes_decipherseed_{decipher_seed}.png"
    return os.path.join(figs_dir, arm, "bif", f"sigma{sigma}", "vspace", name)


def caption(rho: float, r2: float) -> str:
    return f"rho {rho:.3f} · r² {r2:.3f}"


def cell_boxes(n_rows: int, n_cols: int) -> List[List[Box]]:
    """Image+caption box of each grid cell, [row][col], inside the slide below the headers."""
    left0 = MARGIN + ROW_LABEL_W
    top0 = HEADER_TOP + HEADER_H + GAP
    cell_w = (SLIDE_W_IN - MARGIN - left0) / n_cols
    cell_h = (SLIDE_H_IN - MARGIN - top0) / n_rows
    return [
        [(left0 + c * cell_w, top0 + r * cell_h, cell_w, cell_h) for c in range(n_cols)]
        for r in range(n_rows)
    ]


def row_boxes(n: int) -> List[Box]:
    """n equal boxes side by side, filling the slide below the title."""
    top = TITLE_TOP + TITLE_H + GAP
    w = (SLIDE_W_IN - 2 * MARGIN) / n
    return [(MARGIN + i * w, top, w, SLIDE_H_IN - MARGIN - top) for i in range(n)]


def seed_range(values: Sequence[int]) -> str:
    """ "3–5" for a run of consecutive integers, else "1, 3, 4"; sorted either way."""
    v = sorted(set(values))
    if len(v) > 1 and v == list(range(v[0], v[-1] + 1)):
        return f"{v[0]}–{v[-1]}"
    return ", ".join(map(str, v))


def pooled_label(sim_seeds: Sequence[int], decipher_seeds: Sequence[int]) -> str:
    """E.g. "sim seeds 3–5 × decipherseeds 1–6": the runs a compare plot averages over."""
    sim = f"sim seed{'s' if len(set(sim_seeds)) > 1 else ''} {seed_range(sim_seeds)}"
    ds = f"decipherseed{'s' if len(set(decipher_seeds)) > 1 else ''} {seed_range(decipher_seeds)}"
    return f"{sim} × {ds}"


def pooled_seeds(
    ledger: Ledger, models: Sequence[str], sigmas: Sequence[float]
) -> Tuple[List[int], List[int]]:
    """Sorted sim seeds and decipherseeds of the ledger runs on the deck's models and sigmas.

    The compare plots average over exactly these runs (same ledgers, same arm/sigma filter).
    """
    keys = [k for k in ledger if k[0] in models and k[1] in set(sigmas)]
    if not keys:
        raise ValueError(f"no ledger runs for models {list(models)} at sigmas {list(sigmas)}")
    return sorted({k[2] for k in keys}), sorted({k[3] for k in keys})


def compare_slide_title(group_title: str, pooled: str) -> str:
    """`pooled` is a `pooled_label`: compare plots pool sim seeds, unlike the v-space slides."""
    return f"{group_title} | {pooled} | bif, gene-mode"


def unplaced_pngs(png_names: Sequence[str], groups: Groups) -> List[str]:
    """PNG names that no group places on a slide, sorted."""
    placed = {name for _, names in groups for name in names}
    return sorted(set(png_names) - placed)


def fit_image(img_w: float, img_h: float, box: Box) -> Box:
    """Largest rectangle of the image's aspect that fits the box, centred in it."""
    left, top, bw, bh = box
    if img_w <= 0 or img_h <= 0:
        raise ValueError(f"image size must be positive, got {img_w}x{img_h}")
    scale = min(bw / img_w, bh / img_h)
    w, h = img_w * scale, img_h * scale
    return left + (bw - w) / 2, top + (bh - h) / 2, w, h


def cell_inputs(
    figs_dir: str,
    ledger: Ledger,
    sigma: float,
    seed: int,
    models: Sequence[str],
    decipher_seeds: Sequence[int],
) -> List[List[Tuple[str, str]]]:
    """[row][col] -> (png path, caption) for one slide; raises on a missing ledger row."""
    grid = []
    for ds in decipher_seeds:
        row = []
        for m in models:
            key = (m, sigma, seed, ds)
            if key not in ledger:
                raise KeyError(f"no ledger row for {key}")
            row.append((png_path(figs_dir, m, sigma, seed, ds), caption(*ledger[key])))
        grid.append(row)
    return grid


# ---- I/O --------------------------------------------------------------------------------


def read_ledger(paths: Union[str, Sequence[str]]) -> Ledger:
    """(model, sigma, sim seed, decipherseed) -> (rho, r2_overall), gene-mode bif rows only.

    Several ledgers are merged; a key found in more than one of them raises.
    """
    out: Ledger = {}
    for path in [paths] if isinstance(paths, str) else paths:
        if not os.path.isfile(path):
            raise FileNotFoundError(f"ledger not found: {path}")
        seen = {}
        with open(path, newline="") as f:
            for r in csv.DictReader(f):
                if r["bifurcation"] != "bif" or r["batch_mode"] != "genes" or r.get("error"):
                    continue
                key = (
                    r["model"],
                    float(r["shift_sigma"]),
                    int(r["seed"]),
                    int(r["decipher_seed"]),
                )
                seen[key] = (float(r["rho"]), float(r["r2_overall"]))
        dup = sorted(set(seen) & set(out))
        if dup:
            raise ValueError(f"{len(dup)} runs in more than one ledger, e.g. {dup[0]} ({path})")
        out.update(seen)
    return out


def _text(slide, text: str, box: Box, size: int, bold: bool = False) -> None:
    tb = slide.shapes.add_textbox(*(Inches(v) for v in box))
    p = tb.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(0x20, 0x20, 0x20)


def _titled_slide(prs, title: str):
    slide = prs.slides.add_slide(prs.slide_layouts[5])  # title only
    t = slide.shapes.title
    t.text = title
    t.left, t.top = Inches(MARGIN), Inches(TITLE_TOP)
    t.width, t.height = Inches(SLIDE_W_IN - 2 * MARGIN), Inches(TITLE_H)
    t.text_frame.paragraphs[0].runs[0].font.size = Pt(26)
    return slide


def _add_fitted_picture(slide, png: str, box: Box) -> None:
    with Image.open(png) as im:
        iw, ih = im.size
    x, y, pw, ph = fit_image(iw, ih, box)
    slide.shapes.add_picture(png, Inches(x), Inches(y), Inches(pw), Inches(ph))


def _compare_pngs(compare_dir: str, groups: Groups) -> None:
    """Raise if a grouped PNG is missing or a PNG in compare_dir is in no group."""
    if not os.path.isdir(compare_dir):
        raise FileNotFoundError(f"compare dir not found: {compare_dir}")
    present = [n for n in os.listdir(compare_dir) if n.endswith(".png")]
    missing = [n for _, names in groups for n in names if n not in present]
    if missing:
        raise FileNotFoundError(f"missing compare PNGs in {compare_dir}: {missing}")
    extra = unplaced_pngs(present, groups)
    if extra:
        raise ValueError(f"compare PNGs in no COMPARE_GROUPS entry: {extra}")


def build_deck(
    ledger_path: Union[str, Sequence[str]],
    figs_dir: str,
    compare_dir: str,
    out_path: str,
    seed: int = 3,
    sigmas: Sequence[float] = SIGMAS,
    models: Sequence[str] = MODELS,
    decipher_seeds: Sequence[int] = DECIPHER_SEEDS,
    compare_groups: Groups = COMPARE_GROUPS,
) -> None:
    ledger = read_ledger(ledger_path)
    _compare_pngs(compare_dir, compare_groups)
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(SLIDE_W_IN), Inches(SLIDE_H_IN)
    boxes = cell_boxes(len(decipher_seeds), len(models))
    for sigma in sigmas:
        grid = cell_inputs(figs_dir, ledger, sigma, seed, models, decipher_seeds)
        slide = _titled_slide(prs, slide_title(sigma, seed))
        for c, m in enumerate(models):
            left, _, w, _ = boxes[0][c]
            _text(slide, m, (left, HEADER_TOP, w, HEADER_H), 16, bold=True)
        for r, ds in enumerate(decipher_seeds):
            _, top, _, h = boxes[r][0]
            label_box = (MARGIN, top + h / 2 - 0.2, ROW_LABEL_W, 0.4)
            _text(slide, f"decipherseed {ds}", label_box, 13, bold=True)
            for c in range(len(models)):
                png, cap = grid[r][c]
                if not os.path.isfile(png):
                    raise FileNotFoundError(f"missing v-space PNG: {png}")
                left, top, w, h = boxes[r][c]
                _add_fitted_picture(slide, png, (left, top, w, h - CAPTION_H))
                _text(slide, cap, (left, top + h - CAPTION_H, w, CAPTION_H), 11)
    pooled = pooled_label(*pooled_seeds(ledger, models, sigmas))
    for group_title, names in compare_groups:
        slide = _titled_slide(prs, compare_slide_title(group_title, pooled))
        for name, box in zip(names, row_boxes(len(names))):
            _add_fitted_picture(slide, os.path.join(compare_dir, name), box)
    os.makedirs(os.path.dirname(os.path.abspath(out_path)), exist_ok=True)
    prs.save(out_path)


def _sha(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def check_deck(
    deck_path: str,
    figs_dir: str,
    compare_dir: str,
    seed: int = 3,
    sigmas: Sequence[float] = SIGMAS,
    models: Sequence[str] = MODELS,
    decipher_seeds: Sequence[int] = DECIPHER_SEEDS,
    compare_groups: Groups = COMPARE_GROUPS,
    *,
    ledger_path: Union[str, Sequence[str]],
) -> None:
    """Reopen the deck and assert titles, picture counts, picture bytes and bounds.

    `ledger_path` is the build's ledger(s), for the pooled seeds in the compare-slide titles.
    """
    pooled = pooled_label(*pooled_seeds(read_ledger(ledger_path), models, sigmas))
    prs = Presentation(deck_path)
    slides = list(prs.slides)
    # (title, source PNGs in insertion order) per expected slide; v-space pictures are added
    # row-major: decipherseed rows, model columns
    expected = [
        (
            slide_title(sigma, seed),
            [png_path(figs_dir, m, sigma, seed, ds) for ds in decipher_seeds for m in models],
        )
        for sigma in sigmas
    ] + [
        (compare_slide_title(t, pooled), [os.path.join(compare_dir, n) for n in names])
        for t, names in compare_groups
    ]
    assert len(slides) == len(expected), f"{len(slides)} slides, expected {len(expected)}"
    sw, sh = prs.slide_width, prs.slide_height
    for slide, (want, sources) in zip(slides, expected):
        got = slide.shapes.title.text
        assert got == want, f"title {got!r} != {want!r}"
        pics = [s for s in slide.shapes if s.shape_type == 13]  # MSO_SHAPE_TYPE.PICTURE
        assert len(pics) == len(sources), f"{want}: {len(pics)} pictures, expected {len(sources)}"
        for pic, src in zip(pics, sources):
            with open(src, "rb") as f:
                assert _sha(pic.image.blob) == _sha(f.read()), f"{want}: picture != {src}"
            inside = (
                pic.left >= 0
                and pic.top >= 0
                and pic.left + pic.width <= sw
                and pic.top + pic.height <= sh
            )
            assert inside, f"{want}: picture of {src} outside the slide"


def main(argv: Optional[Sequence[str]] = None, compare_groups: Groups = COMPARE_GROUPS) -> None:
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("--ledger", action="append", help=f"repeatable (default {DEFAULT_LEDGER})")
    ap.add_argument("--figs-dir", default=DEFAULT_FIGS_DIR)
    ap.add_argument("--compare-dir", default=DEFAULT_COMPARE_DIR)
    ap.add_argument("--out", default=DEFAULT_OUT)
    ap.add_argument("--seed", type=int, default=3)
    ap.add_argument("--models", nargs="+", default=MODELS)
    ap.add_argument("--sigmas", type=float, nargs="+", default=SIGMAS)
    ap.add_argument("--decipher-seeds", type=int, nargs="+", default=DECIPHER_SEEDS)
    ap.add_argument("--check", action="store_true", help="only check an existing --out deck")
    args = ap.parse_args(argv)
    grid = dict(
        sigmas=args.sigmas,
        models=args.models,
        decipher_seeds=args.decipher_seeds,
        compare_groups=compare_groups,
    )
    ledgers = args.ledger or [DEFAULT_LEDGER]
    if not args.check:
        build_deck(ledgers, args.figs_dir, args.compare_dir, args.out, args.seed, **grid)
        print(f"built {args.out}")
    check_deck(args.out, args.figs_dir, args.compare_dir, args.seed, **grid, ledger_path=ledgers)
    print("check OK")


if __name__ == "__main__":
    main()
