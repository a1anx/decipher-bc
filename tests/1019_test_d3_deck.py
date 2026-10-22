"""No-training tests for the d3 deck builder, on a tiny fake figs tree."""

import csv
import importlib.util
import os

import pytest
from PIL import Image

_SRC = os.path.join(
    os.path.dirname(__file__),
    "..",
    "Simulated Data",
    "Simulated Data Sweep Pipeline",
    "1019_build_d3_deck.py",
)
_spec = importlib.util.spec_from_file_location("build_d3_deck", _SRC)
deck = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(deck)

SIGMAS = [0.5, 2.0]
MODELS = ["native", "model5"]
DSEEDS = [1, 2]
GROUPS = [
    ("Rho", ["rho_vs_sigma.png", "rho_plus_vs_sigma.png"]),
    ("Epochs", ["n_epochs_vs_sigma.png"]),
]
KW = dict(sigmas=SIGMAS, models=MODELS, decipher_seeds=DSEEDS, compare_groups=GROUPS)


@pytest.fixture
def tree(tmp_path):
    figs = tmp_path / "figs"
    rows = []
    for s in SIGMAS:
        for m in MODELS:
            for ds in DSEEDS:
                p = deck.png_path(str(figs), m, s, 3, ds)
                os.makedirs(os.path.dirname(p), exist_ok=True)
                shade = int(10 * s) + 40 * ds + (7 if m == "native" else 0)
                Image.new("RGB", (40, 30), (shade, 0, 0)).save(p)
                rows.append(
                    dict(
                        model=m,
                        batch_mode="genes",
                        bifurcation="bif",
                        shift_sigma=s,
                        seed=3,
                        decipher_seed=ds,
                        rho=0.9,
                        r2_overall=0.8,
                        error="",
                    )
                )
    compare = tmp_path / "compare"
    compare.mkdir()
    for i, (_, names) in enumerate(GROUPS):
        for j, name in enumerate(names):
            Image.new("RGB", (70, 51), (0, 30 * i + 10 * j, 0)).save(compare / name)
    ledger = tmp_path / "ledger.csv"
    with open(ledger, "w", newline="") as f:
        w = csv.DictWriter(f, fieldnames=list(rows[0]))
        w.writeheader()
        w.writerows(rows)
    return str(ledger), str(figs), str(compare), str(tmp_path / "out.pptx")


def test_build_and_check_pass(tree):
    ledger, figs, compare, out = tree
    deck.build_deck(ledger, figs, compare, out, **KW)
    deck.check_deck(out, figs, compare, ledger_path=ledger, **KW)


def test_missing_png_raises_with_path(tree):
    ledger, figs, compare, out = tree
    os.remove(deck.png_path(figs, "model5", 2.0, 3, 2))
    with pytest.raises(FileNotFoundError, match="vspace_sigma2_seed3_model5"):
        deck.build_deck(ledger, figs, compare, out, **KW)


def test_missing_ledger_row_raises(tree):
    ledger, figs, compare, out = tree
    with pytest.raises(KeyError, match="model7"):
        deck.build_deck(ledger, figs, compare, out, **{**KW, "models": ["model7"]})


def test_tampered_picture_fails_check(tree):
    ledger, figs, compare, out = tree
    deck.build_deck(ledger, figs, compare, out, **KW)
    src = deck.png_path(figs, "native", 0.5, 3, 1)
    Image.new("RGB", (40, 30), (255, 255, 255)).save(src)  # source changed after the build
    with pytest.raises(AssertionError, match="picture != "):
        deck.check_deck(out, figs, compare, ledger_path=ledger, **KW)


def test_title_path_and_caption_text():
    assert deck.slide_title(0.5, 3) == "σ = 0.5 | sim seed 3 | bif, gene-mode"
    assert deck.slide_title(10.0, 3) == "σ = 10 | sim seed 3 | bif, gene-mode"
    assert deck.png_path("F", "native", 2.0, 3, 1) == os.path.join(
        "F",
        "native_genes",
        "bif",
        "sigma2.0",
        "vspace",
        "vspace_sigma2_seed3_native_bif_genes_decipherseed_1.png",
    )
    assert deck.caption(0.94142, 0.65677) == "rho 0.941 · r² 0.657"


@pytest.mark.parametrize("size", [(1000, 10), (10, 1000), (30, 40), (1, 1)])
def test_fit_keeps_aspect_and_stays_in_box(size):
    box = (2.0, 1.0, 3.0, 2.0)
    x, y, w, h = deck.fit_image(*size, box)
    assert w / h == pytest.approx(size[0] / size[1])
    assert box[0] - 1e-9 <= x and x + w <= box[0] + box[2] + 1e-9
    assert box[1] - 1e-9 <= y and y + h <= box[1] + box[3] + 1e-9
    assert w == pytest.approx(box[2]) or h == pytest.approx(box[3])  # fills one dimension


def test_full_grid_cells_inside_slide():
    for row in deck.cell_boxes(3, 5):
        for left, top, w, h in row:
            assert left >= 0 and top >= 0
            assert left + w <= deck.SLIDE_W_IN + 1e-9 and top + h <= deck.SLIDE_H_IN + 1e-9


def _edit_first_slide(out, edit):
    from pptx import Presentation

    prs = Presentation(out)
    edit(prs, list(prs.slides[0].shapes))
    prs.save(out)


def test_picture_outside_slide_fails_check(tree):
    ledger, figs, compare, out = tree
    deck.build_deck(ledger, figs, compare, out, **KW)

    def push_off(prs, shapes):
        pic = next(s for s in shapes if s.shape_type == 13)
        pic.left = prs.slide_width - pic.width // 2

    _edit_first_slide(out, push_off)
    with pytest.raises(AssertionError, match="outside the slide"):
        deck.check_deck(out, figs, compare, ledger_path=ledger, **KW)


def test_missing_picture_fails_check(tree):
    ledger, figs, compare, out = tree
    deck.build_deck(ledger, figs, compare, out, **KW)

    def drop(prs, shapes):
        pic = next(s for s in shapes if s.shape_type == 13)
        pic._element.getparent().remove(pic._element)

    _edit_first_slide(out, drop)
    with pytest.raises(AssertionError, match="3 pictures, expected 4"):
        deck.check_deck(out, figs, compare, ledger_path=ledger, **KW)


def test_compare_slides_follow_sigma_slides_in_group_order(tree):
    ledger, figs, compare, out = tree
    deck.build_deck(ledger, figs, compare, out, **KW)
    from pptx import Presentation

    titles = [s.shapes.title.text for s in Presentation(out).slides]
    assert titles == [deck.slide_title(s, 3) for s in SIGMAS] + [
        "Rho | sim seed 3 × decipherseeds 1–2 | bif, gene-mode",
        "Epochs | sim seed 3 × decipherseeds 1–2 | bif, gene-mode",
    ]


def test_missing_compare_png_raises_with_name(tree):
    ledger, figs, compare, out = tree
    os.remove(os.path.join(compare, "rho_plus_vs_sigma.png"))
    with pytest.raises(FileNotFoundError, match="rho_plus_vs_sigma.png"):
        deck.build_deck(ledger, figs, compare, out, **KW)


def test_ungrouped_compare_png_raises(tree):
    ledger, figs, compare, out = tree
    Image.new("RGB", (5, 5)).save(os.path.join(compare, "new_metric_vs_sigma.png"))
    with pytest.raises(ValueError, match="new_metric_vs_sigma.png"):
        deck.build_deck(ledger, figs, compare, out, **KW)


def test_tampered_compare_picture_fails_check(tree):
    ledger, figs, compare, out = tree
    deck.build_deck(ledger, figs, compare, out, **KW)
    Image.new("RGB", (70, 51), (255, 255, 255)).save(os.path.join(compare, "n_epochs_vs_sigma.png"))
    with pytest.raises(AssertionError, match="Epochs.*picture != "):
        deck.check_deck(out, figs, compare, ledger_path=ledger, **KW)


def test_default_groups_place_each_compare_png_once():
    names = [n for _, group in deck.COMPARE_GROUPS for n in group]
    assert sorted(names) == sorted(
        f"{m}_vs_sigma.png"
        for m in [
            "rho",
            "rho_plus",
            "rho_minus",
            "r2_overall",
            "r2_per_gene_median",
            "branch_asw",
            "batch_asw_v",
            "best_epoch",
            "n_epochs",
        ]
    )


@pytest.mark.parametrize("n", [1, 2, 3])
def test_row_boxes_tile_slide_width_below_title(n):
    boxes = deck.row_boxes(n)
    assert boxes[0][0] == pytest.approx(deck.MARGIN)
    assert boxes[-1][0] + boxes[-1][2] == pytest.approx(deck.SLIDE_W_IN - deck.MARGIN)
    for (l1, _, w1, _), (l2, _, _, _) in zip(boxes, boxes[1:]):
        assert l1 + w1 == pytest.approx(l2)
    for _, top, _, h in boxes:
        assert top >= deck.TITLE_TOP + deck.TITLE_H
        assert top + h == pytest.approx(deck.SLIDE_H_IN - deck.MARGIN)


def test_unplaced_pngs_lists_only_ungrouped_sorted():
    assert deck.unplaced_pngs(["b.png", "a.png", "c.png"], [("G", ["c.png"])]) == ["a.png", "b.png"]
    assert deck.unplaced_pngs([], [("G", ["c.png"])]) == []


def _split_ledger(ledger: str, tmp_path) -> list:
    """Split the fixture ledger into two files by decipherseed (1 | 2)."""
    with open(ledger, newline="") as f:
        rows = list(csv.DictReader(f))
    paths = []
    for ds in ("1", "2"):
        p = tmp_path / f"ledger_ds{ds}.csv"
        with open(p, "w", newline="") as f:
            w = csv.DictWriter(f, fieldnames=list(rows[0]))
            w.writeheader()
            w.writerows(r for r in rows if r["decipher_seed"] == ds)
        paths.append(str(p))
    return paths


def test_read_ledger_merges_several_files(tree, tmp_path):
    ledger = tree[0]

    merged = deck.read_ledger(_split_ledger(ledger, tmp_path))

    assert merged == deck.read_ledger(ledger)


def test_read_ledger_raises_on_key_in_two_files(tree):
    ledger = tree[0]
    with pytest.raises(ValueError, match="in more than one ledger"):
        deck.read_ledger([ledger, ledger])


def test_main_builds_and_checks_deck_from_two_ledgers(tree, tmp_path):
    ledger, figs, compare, out = tree
    a, b = _split_ledger(ledger, tmp_path)
    argv = ["--ledger", a, "--ledger", b, "--figs-dir", figs, "--compare-dir", compare]
    argv += ["--out", out, "--models", *MODELS, "--sigmas", *map(str, SIGMAS)]
    argv += ["--decipher-seeds", *map(str, DSEEDS)]

    deck.main(argv, compare_groups=GROUPS)

    deck.check_deck(out, figs, compare, ledger_path=ledger, **KW)


@pytest.mark.parametrize(
    "values, text", [([3, 4, 5], "3–5"), ([5, 3, 4], "3–5"), ([3], "3"), ([1, 3, 4], "1, 3, 4")]
)
def test_seed_range_text(values, text):
    assert deck.seed_range(values) == text


@pytest.mark.parametrize(
    "sim, ds, label",
    [
        ([3, 4, 5], [1, 2, 3, 4, 5, 6], "sim seeds 3–5 × decipherseeds 1–6"),
        ([3], [1], "sim seed 3 × decipherseed 1"),
    ],
)
def test_pooled_label_names_both_seed_kinds(sim, ds, label):
    assert deck.pooled_label(sim, ds) == label


def test_pooled_seeds_counts_only_the_decks_models_and_sigmas():
    ledger = {
        ("model5", 7.5, 3, 4): (0.9, 0.9),
        ("model5", 10.0, 5, 6): (0.9, 0.9),
        ("model5", 2.0, 4, 9): (0.9, 0.9),  # sigma not on the deck
        ("native", 7.5, 6, 1): (0.9, 0.9),  # model not on the deck
    }

    assert deck.pooled_seeds(ledger, ["model5"], [7.5, 10.0]) == ([3, 5], [4, 6])


def test_pooled_seeds_raises_when_no_run_matches():
    with pytest.raises(ValueError, match="no ledger runs"):
        deck.pooled_seeds({("model5", 7.5, 3, 4): (0.9, 0.9)}, ["model7"], [7.5])
